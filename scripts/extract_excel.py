#!/usr/bin/env python3
"""
Extrai as bases de dados da planilha do Treinador (FG_Planilha.xlsx) para os
ficheiros que alimentam o Supabase:

  supabase/migrations/20260916T0005_gfit_seed_exercises.sql
  supabase/migrations/20260916T0006_gfit_seed_foods.sql
  supabase/seed/exercises.json
  supabase/seed/foods.json

O SQL é o artefacto canónico (corre com o resto das migrações); o JSON é a mesma
informação em formato portátil, usada no arranque para carregar os dados sem ter
de colar 130 KB de SQL.

    pip install openpyxl && python3 scripts/extract_excel.py
"""
import json
import re
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "FG_Planilha.xlsx"

# Pesos de contributo de volume por coluna de músculo, como na planilha:
# as 3 primeiras colunas contam série inteira, as 2 seguintes meia, a última 0.3.
MUSCLE_WEIGHTS = [1, 1, 1, 0.5, 0.5, 0.3]

EQUIPMENT_RULES = [
    (r"halter", "Halteres"),
    (r"\bbarra\b|\bz\b|smith", "Barra"),
    (r"cabo|cross ?over|polia", "Cabos"),
    (r"m[aá]quina|machine|hack|leg press|peck deck|scott", "Máquina"),
    (r"el[aá]stico|band", "Elástico"),
    (r"\btrx\b|suspens", "TRX"),
    (r"kettlebell", "Kettlebell"),
    (r"bola|fitball|swiss", "Bola"),
    (r"\bbanco\b", "Banco"),
]


def deaccent(value: str) -> str:
    return unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()


def slugify(value: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", deaccent(value)).strip("_").lower()


def sql_str(value) -> str:
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def sql_num(value) -> str:
    number = float(value)
    return str(int(number)) if number == int(number) else str(number)


def equipment_for(name: str) -> str:
    haystack = deaccent(name).lower()
    for pattern, label in EQUIPMENT_RULES:
        if re.search(pattern, haystack):
            return label
    return "Livre"


def clean(value):
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def read_muscles(sheet):
    """A lista-mestra de músculos vive na coluna L da folha de exercícios."""
    names = [clean(sheet.cell(row=r, column=12).value) for r in range(3, 19)]
    return [
        {"slug": slugify(name), "name": name, "sort_order": index * 10}
        for index, name in enumerate(n for n in names if n)
    ]


def read_exercises(sheet, muscles):
    by_name = {m["name"].lower(): m["slug"] for m in muscles}
    exercises, seen = [], set()

    for row in range(3, sheet.max_row + 1):
        name = clean(sheet.cell(row=row, column=3).value)
        if not name or name.lower() in seen:
            continue
        seen.add(name.lower())

        video = clean(sheet.cell(row=row, column=4).value)
        if video and not video.startswith("http"):
            video = None

        targeted = []
        for index, column in enumerate(range(5, 11)):
            muscle = clean(sheet.cell(row=row, column=column).value)
            slug = by_name.get(muscle.lower()) if muscle else None
            if slug and all(m["muscle"] != slug for m in targeted):
                targeted.append({"muscle": slug, "weight": MUSCLE_WEIGHTS[index]})

        exercises.append(
            {
                "name": name,
                "pattern": clean(sheet.cell(row=row, column=1).value),
                "category": clean(sheet.cell(row=row, column=2).value),
                "video_url": video,
                "primary_muscle": targeted[0]["muscle"] if targeted else None,
                "muscles": targeted,
                "equipment": equipment_for(name),
            }
        )
    return exercises


def read_foods(sheet):
    foods, seen = [], set()

    for row in range(2, sheet.max_row + 1):
        name = clean(sheet.cell(row=row, column=1).value)
        if not name or name.lower() in seen:
            continue
        macro_cells = [sheet.cell(row=row, column=c).value for c in (4, 5, 6)]
        if all(cell is None for cell in macro_cells):
            continue
        seen.add(name.lower())

        def number(column):
            try:
                return round(float(sheet.cell(row=row, column=column).value), 2)
            except (TypeError, ValueError):
                return 0.0

        foods.append(
            {
                "name": name,
                "base_qty": number(2) or 100.0,
                "unit": (clean(sheet.cell(row=row, column=3).value) or "g")[:12],
                "protein_g": number(4),
                "fat_g": number(5),
                "carb_g": number(6),
            }
        )
    return foods


def write_exercises_sql(path, muscles, exercises):
    lines = [
        "-- GFit — músculos e biblioteca de exercícios.",
        "-- Gerado por scripts/extract_excel.py a partir de FG_Planilha.xlsx.",
        "",
        "insert into gfit.muscles (slug, name, sort_order) values",
    ]
    lines.append(
        ",\n".join(
            f"  ({sql_str(m['slug'])},{sql_str(m['name'])},{m['sort_order']})"
            for m in muscles
        )
        + "\non conflict (slug) do nothing;\n"
    )
    lines.append(
        "-- Os músculos viajam como 'slug:peso|slug:peso' e viram jsonb à chegada.\n"
        "insert into gfit.exercises\n"
        "  (name, pattern, category, video_url, primary_muscle, muscles, equipment)\n"
        "select d.name, d.pattern, d.category, d.video_url,\n"
        "       nullif(split_part(split_part(d.m, '|', 1), ':', 1), ''),\n"
        "       coalesce((\n"
        "         select jsonb_agg(jsonb_build_object(\n"
        "                  'muscle', split_part(part, ':', 1),\n"
        "                  'weight', split_part(part, ':', 2)::numeric))\n"
        "           from unnest(string_to_array(d.m, '|')) part\n"
        "          where part <> ''\n"
        "       ), '[]'::jsonb),\n"
        "       d.equipment\n"
        "from (values"
    )
    rows = []
    for ex in exercises:
        packed = "|".join(f"{m['muscle']}:{m['weight']}" for m in ex["muscles"])
        rows.append(
            f"  ({sql_str(ex['name'])},{sql_str(ex['pattern'])},"
            f"{sql_str(ex['category'])},{sql_str(ex['video_url'])},"
            f"{sql_str(packed)},{sql_str(ex['equipment'])})"
        )
    lines.append(",\n".join(rows))
    lines.append(") as d(name, pattern, category, video_url, m, equipment);")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_foods_sql(path, foods):
    per_100g = [f for f in foods if f["base_qty"] == 100.0 and f["unit"] == "g"]
    others = [f for f in foods if f not in per_100g]

    lines = [
        "-- GFit — base de alimentos.",
        "-- Gerado por scripts/extract_excel.py a partir de FG_Planilha.xlsx.",
        "",
        "-- Valores por 100 g.",
        "insert into gfit.foods (name, protein_g, fat_g, carb_g) values",
    ]
    lines.append(
        ",\n".join(
            f"  ({sql_str(f['name'])},{sql_num(f['protein_g'])},"
            f"{sql_num(f['fat_g'])},{sql_num(f['carb_g'])})"
            for f in per_100g
        )
        + ";"
    )
    if others:
        lines += [
            "",
            "-- Outras quantidades/unidades.",
            "insert into gfit.foods (name, base_qty, unit, protein_g, fat_g, carb_g) values",
            ",\n".join(
                f"  ({sql_str(f['name'])},{sql_num(f['base_qty'])},{sql_str(f['unit'])},"
                f"{sql_num(f['protein_g'])},{sql_num(f['fat_g'])},{sql_num(f['carb_g'])})"
                for f in others
            )
            + ";",
        ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    workbook = openpyxl.load_workbook(XLSX, data_only=True)
    exercise_sheet = workbook["Base dados Exercicios"]

    muscles = read_muscles(exercise_sheet)
    exercises = read_exercises(exercise_sheet, muscles)
    foods = read_foods(workbook["Base dados dieta"])

    migrations = ROOT / "supabase" / "migrations"
    seed = ROOT / "supabase" / "seed"
    migrations.mkdir(parents=True, exist_ok=True)
    seed.mkdir(parents=True, exist_ok=True)

    write_exercises_sql(migrations / "20260916T0005_gfit_seed_exercises.sql", muscles, exercises)
    write_foods_sql(migrations / "20260916T0006_gfit_seed_foods.sql", foods)

    (seed / "exercises.json").write_text(
        json.dumps({"muscles": muscles, "exercises": exercises}, ensure_ascii=False),
        encoding="utf-8",
    )
    (seed / "foods.json").write_text(
        json.dumps(foods, ensure_ascii=False), encoding="utf-8"
    )

    with_video = sum(1 for e in exercises if e["video_url"])
    print(f"{len(muscles)} músculos")
    print(f"{len(exercises)} exercícios ({with_video} com vídeo)")
    print(f"{len(foods)} alimentos")


if __name__ == "__main__":
    main()
